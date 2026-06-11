import os, re
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Emu

DECK_DIR = os.path.dirname(os.path.abspath(__file__))
HTML_FILE = os.path.join(DECK_DIR, "VeraCare_Deck.html")
SCREENSHOT_DIR = os.path.join(DECK_DIR, "slide_screenshots")
OUTPUT_PPTX = os.path.join(DECK_DIR, "VeraCare_Pitch_Deck.pptx")
SLIDE_W, SLIDE_H = 1920, 1080
SCALE = 2  # render at 2x for crisp screenshots

os.makedirs(SCREENSHOT_DIR, exist_ok=True)

with open(HTML_FILE, "r") as f:
    html = f.read()

head_match = re.search(r"(<head.*?</head>)", html, re.DOTALL)
head_block = head_match.group(1) if head_match else "<head></head>"

slides = re.split(r"(?=<div class=\"slide\")", html)
slide_blocks = [s for s in slides if s.strip().startswith('<div class="slide"')]

print(f"Found {len(slide_blocks)} slides")

with sync_playwright() as p:
    browser = p.chromium.launch()
    page = browser.new_page(
        viewport={"width": SLIDE_W, "height": SLIDE_H},
        device_scale_factor=SCALE,
    )

    for i, slide_html in enumerate(slide_blocks):
        single_html = f"""<!DOCTYPE html>
<html lang="en">
{head_block}
<style>
html,body{{margin:0!important;padding:0!important;overflow:hidden!important;background:#0A0F1C!important;width:100vw!important;height:100vh!important}}
.slide{{width:100vw!important;height:100vh!important}}
</style>
<body>
{slide_html}
</body>
</html>"""

        tmp_path = os.path.join(DECK_DIR, f"_tmp_slide_{i+1}.html")
        with open(tmp_path, "w") as tf:
            tf.write(single_html)

        out_path = os.path.join(SCREENSHOT_DIR, f"slide_{i+1:02d}.png")
        page.goto(f"file://{tmp_path}", wait_until="networkidle")
        page.wait_for_timeout(400)
        page.screenshot(path=out_path, clip={"x": 0, "y": 0, "width": SLIDE_W, "height": SLIDE_H})
        os.remove(tmp_path)

        if os.path.exists(out_path) and os.path.getsize(out_path) > 1000:
            print(f"  Captured slide {i+1}")
        else:
            print(f"  WARNING: slide {i+1} may have failed")

    browser.close()

print("\nBuilding PPTX...")

prs = Presentation()
prs.slide_width = Emu(SLIDE_W * 914400 // 96)
prs.slide_height = Emu(SLIDE_H * 914400 // 96)

blank_layout = prs.slide_layouts[6]

for i in range(len(slide_blocks)):
    img_path = os.path.join(SCREENSHOT_DIR, f"slide_{i+1:02d}.png")
    if not os.path.exists(img_path):
        print(f"  Skipping slide {i+1} - no screenshot")
        continue
    slide = prs.slides.add_slide(blank_layout)
    slide.shapes.add_picture(img_path, Emu(0), Emu(0), prs.slide_width, prs.slide_height)
    print(f"  Added slide {i+1}")

prs.save(OUTPUT_PPTX)
print(f"\nDone! PPTX saved to: {OUTPUT_PPTX}")
